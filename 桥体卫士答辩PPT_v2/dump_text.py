# -*- coding: utf-8 -*-
"""dump_text.py —— 提取 PPTX 每页可见文本 + notes，生成比对用清单。

用法:
    <bundled-python> dump_text.py [deck.pptx] [--out _work/pptx_text.txt]
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent


def slide_texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.shape_type == 13:
            out.append(("PICTURE", "[[位图]]"))
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(("text", sh.text_frame.text))
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                out.append(("table", " | ".join(c.text for c in r.cells)))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck", nargs="?", default=str(HERE / "桥体卫士答辩PPT_12分钟版.pptx"))
    ap.add_argument("--out", default=str(HERE / "_work" / "pptx_text.txt"))
    args = ap.parse_args()

    prs = Presentation(args.deck)
    spec = json.loads((HERE / "_spec" / "SPEC_content.json").read_text(encoding="utf-8"))
    by_id = {s["slide_id"]: s for s in spec["slides"]}

    L = []
    L.append(f"# PPTX 可见文本提取\n# 文件: {Path(args.deck).name}\n# 页数: {len(prs.slides)}\n")
    for i, slide in enumerate(prs.slides, 1):
        sid = spec["slides"][i - 1]["slide_id"] if i <= len(spec["slides"]) else f"#{i}"
        L.append("=" * 78)
        L.append(f"## 第 {i:02d} 页  {sid}   标题: {by_id.get(sid, {}).get('title')!r}")
        L.append("-" * 78)
        for kind, txt in slide_texts(slide):
            L.append(f"[{kind}] {txt}")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        L.append(f"[notes] {notes}")
        L.append("")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(L), encoding="utf-8")
    print(f"已写出: {out}  ({out.stat().st_size} bytes)")
    print(f"页数: {len(prs.slides)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
