# -*- coding: utf-8 -*-
"""probe_timing.py —— 用 python-pptx 逐页检查 p:timing 是否存在（排查误报）。"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

p = sys.argv[1]
prs = Presentation(p)
print("页数:", len(prs.slides))
hits = 0
for i, s in enumerate(prs.slides, 1):
    t = s._element.find(qn("p:timing"))
    if t is not None:
        hits += 1
        print(f"  第{i}页 有 p:timing, 子节点数={len(t)}")
print("有 timing 的页数:", hits)
