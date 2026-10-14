# -*- coding: utf-8 -*-
"""诊断渲染偏移：对比规格坐标与 pptx 内实际坐标，并 dump 一个形状的 XML。"""
import pathlib
import sys

from pptx import Presentation
from pptx.util import Emu

p = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "桥体卫士答辩PPT_v2/_work/smoke/smoke.pptx")
prs = Presentation(str(p))
print("slide size EMU:", prs.slide_width, prs.slide_height)
print("slide size in :", prs.slide_width / 914400, "x", prs.slide_height / 914400)
for i, s in enumerate(prs.slides, 1):
    print(f"--- slide {i} ({len(s.shapes)} shapes)")
    for sh in s.shapes:
        txt = (sh.text_frame.text[:24].replace("\n", "|") if sh.has_text_frame else "")
        print(
            f"    {str(sh.shape_type):18s} x={sh.left} y={sh.top} w={sh.width} h={sh.height}  {txt!r}"
        )
    if i == 1:
        for sh in s.shapes:
            if sh.has_text_frame and "桥体卫士" in sh.text_frame.text:
                print("=== title XML ===")
                print(sh._element.xml[:2600])
                break
