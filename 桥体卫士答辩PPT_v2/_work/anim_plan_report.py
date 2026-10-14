# -*- coding: utf-8 -*-
"""anim_plan_report.py —— 打印每页的动画编排（人读），用于确认"故事线"是否合理。

用法: <python> _work/anim_plan_report.py [deck.pptx]
"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402

from build_animated import decoration_ids_for  # noqa: E402
from deck_style import build_cover_anim_groups, build_page_groups  # noqa: E402


def label(sh):
    if sh.has_text_frame and sh.text_frame.text.strip():
        return sh.text_frame.text.strip().replace("\n", " ")[:16]
    return "图形"


def main():
    p = sys.argv[1] if len(sys.argv) > 1 else str(HERE / "桥体卫士答辩PPT_12分钟版.pptx")
    prs = Presentation(p)
    for i, s in enumerate(prs.slides, 1):
        if i == 1:
            gs = build_cover_anim_groups(s, set())
        elif i == 2:
            gs = []
        else:
            gs = build_page_groups(s, decoration_ids_for(s))
        if not gs:
            print(f"第{i:>2}页  仅转场（无元素动画）")
            continue
        chunks = []
        for gi, click in enumerate(gs, 1):
            body = " + ".join(f"{kind}×{len(shapes)}({','.join(label(x) for x in shapes[:3])}…)"
                              for kind, shapes in click)
            chunks.append(f"点击{gi}: {body}")
        print(f"第{i:>2}页  " + "    ||    ".join(chunks))


if __name__ == "__main__":
    main()
