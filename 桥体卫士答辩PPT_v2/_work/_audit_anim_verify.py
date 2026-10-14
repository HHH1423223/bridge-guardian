# -*- coding: utf-8 -*-
"""Read-only independent audit of animation timing in the deck.

Measures ONLY from the PPTX zip bytes. Never writes to any pptx.
"""
import hashlib
import os
import re
import sys
import zipfile
from lxml import etree
from pptx import Presentation

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
Q = lambda t: "{%s}%s" % (P, t)
QA = lambda t: "{%s}%s" % (A, t)

BASE = r"C:\Users\CX\Documents\GitHub\bridge guardian\桥体卫士答辩PPT_v2"
FILES = [
    ("deliverable", os.path.join(BASE, "\u6865\u4f53\u536b\u58eb\u7b54\u8fa9PPT_12\u5206\u949f\u7248.pptx")),
    ("preanim", os.path.join(BASE, "_work", "\u6865\u4f53\u536b\u58eb\u7b54\u8fa9PPT_12\u5206\u949f\u7248.preanim.pptx")),
    ("v1perline", os.path.join(BASE, "_work", "\u6865\u4f53\u536b\u58eb\u7b54\u8fa9PPT_12\u5206\u949f\u7248.v1-perline.pptx")),
]

out = []
def w(*a):
    s = " ".join(str(x) for x in a)
    out.append(s)
    print(s)


def sha(b):
    return hashlib.sha256(b).hexdigest()


def slide_parts(zf):
    names = [n for n in zf.namelist() if re.match(r"^ppt/slides/slide\d+\.xml$", n)]
    return sorted(names, key=lambda n: int(re.search(r"(\d+)", n.split("/")[-1]).group(1)))


def notes_parts(zf):
    names = [n for n in zf.namelist() if re.match(r"^ppt/notesSlides/notesSlide\d+\.xml$", n)]
    return sorted(names, key=lambda n: int(re.search(r"(\d+)", n.split("/")[-1]).group(1)))


def elem_depth(e):
    d = 0
    while e is not None:
        d += 1
        e = e.getparent()
    return d


def max_depth_of(root):
    best = (0, None)
    for e in root.iter():
        if not isinstance(e.tag, str):
            continue
        d = elem_depth(e)
        if d > best[0]:
            best = (d, e.tag)
    return best


def local(tag):
    return tag.split("}")[-1] if isinstance(tag, str) else str(tag)


def strip_anim_bytes(raw):
    s = raw.decode("utf-8")
    tr = re.findall(r"<p:transition\b[^>]*?/>|<p:transition\b.*?</p:transition>", s, flags=re.S)
    s = re.sub(r"<p:transition\b[^>]*?/>|<p:transition\b.*?</p:transition>", "", s, flags=re.S)
    tm = re.findall(r"<p:timing\b.*?</p:timing>", s, flags=re.S)
    s = re.sub(r"<p:timing\b.*?</p:timing>", "", s, flags=re.S)
    return s.encode("utf-8"), len(tr), len(tm)


# ---------------------------------------------------------------- per part dive
def analyze(path, label):
    w("")
    w("=" * 100)
    w("FILE[%s] %s" % (label, path))
    w("  exists=%s size=%d sha256=%s" % (os.path.exists(path), os.path.getsize(path),
                                         sha(open(path, "rb").read())))
    zf = zipfile.ZipFile(path)
    bad = zf.testzip()
    w("  zip testzip (None == ok) = %r ; entries=%d" % (bad, len(zf.namelist())))

    # ---- every XML part parses + depth
    worst = []
    parse_errors = []
    for n in zf.namelist():
        if not (n.endswith(".xml") or n.endswith(".rels")):
            continue
        data = zf.read(n)
        try:
            root = etree.fromstring(data)
        except Exception as ex:
            parse_errors.append((n, repr(ex)))
            continue
        d, tag = max_depth_of(root)
        worst.append((d, n, local(tag)))
    worst.sort(reverse=True)
    w("  XML parts parsed=%d parse_errors=%d" % (len(worst), len(parse_errors)))
    for n, ex in parse_errors:
        w("    !! PARSE ERROR %s %s" % (n, ex))
    w("  deepest XML parts (depth, part, deepest-element):")
    for d, n, t in worst[:6]:
        w("     %4d  %-45s %s %s" % (d, n, t, "  <<< OVER 64" if d > 64 else ""))

    # ---- python-pptx open
    try:
        prs = Presentation(path)
        w("  python-pptx: Presentation() OK  slides=%d  sldSz=%d x %d EMU (%.4f x %.4f in)"
          % (len(prs.slides), prs.slide_width, prs.slide_height,
             prs.slide_width / 914400, prs.slide_height / 914400))
    except Exception as ex:
        w("  python-pptx: OPEN FAILED %r" % (ex,))
        prs = None

    parts = slide_parts(zf)
    rows = []
    detail = {}
    for n in parts:
        idx = int(re.search(r"(\d+)", n.split("/")[-1]).group(1))
        raw = zf.read(n)
        root = etree.fromstring(raw)
        # cNvPr id -> name (whole spTree, incl. nested groups)
        idname = {}
        for c in root.iter(Q("cNvPr")):
            idname[c.get("id")] = c.get("name")
        # child order of p:sld
        order = [local(c.tag) for c in root if isinstance(c.tag, str)]
        # transition
        trs = root.findall(Q("transition"))
        tr_children = [local(c.tag) for t in trs for c in t]
        tr_attrs = {k.split("}")[-1]: v for k, v in trs[0].attrib.items()} if trs else {}
        # timing
        tim = root.find(Q("timing"))
        clicks = []
        ctn_ids_all = [c.get("id") for c in root.iter(Q("cTn"))]
        dup_ctn = sorted({i for i in ctn_ids_all if ctn_ids_all.count(i) > 1 and i is not None})
        if tim is not None:
            mainseq = tim.find("./%s/%s/%s[@nodeType='tmRoot']/%s/%s/%s[@nodeType='mainSeq']/%s"
                               % (Q("tnLst"), Q("par"), Q("cTn"), Q("childTnLst"),
                                  Q("seq"), Q("cTn"), Q("childTnLst")))
            if mainseq is None:
                # path variant: some writers omit nodeType on tmRoot
                mainseq = tim.find("./%s/%s/%s/%s/%s/%s/%s"
                                   % (Q("tnLst"), Q("par"), Q("cTn"), Q("childTnLst"),
                                      Q("seq"), Q("cTn"), Q("childTnLst")))
            if mainseq is not None:
                for gp in mainseq.findall(Q("par")):
                    ctn = gp.find(Q("cTn"))
                    cond = ctn.find("./%s/%s" % (Q("stCondLst"), Q("cond"))) if ctn is not None else None
                    conds = ctn.findall("./%s/%s" % (Q("stCondLst"), Q("cond"))) if ctn is not None else []
                    node_types = {}
                    spids = []
                    effect_ctns = []
                    for c in gp.iter(Q("cTn")):
                        nt = c.get("nodeType")
                        if nt:
                            node_types[nt] = node_types.get(nt, 0) + 1
                        if c.get("presetClass") or c.get("presetID"):
                            effect_ctns.append(c)
                    for sp in gp.iter(Q("spTgt")):
                        spids.append(sp.get("spid"))
                    distinct = []
                    for s in spids:
                        if s not in distinct:
                            distinct.append(s)
                    clicks.append(dict(
                        cond=[(c.get("delay"), c.get("evt")) for c in conds],
                        node_types=node_types,
                        n_effects=len(effect_ctns),
                        spids=spids,
                        distinct=distinct,
                        names=[idname.get(s, "??MISSING:%s" % s) for s in distinct],
                        preset_classes=sorted({c.get("presetClass") for c in effect_ctns}),
                        preset_ids=sorted({c.get("presetID") for c in effect_ctns}),
                        presets_without_set=[c.get("id") for c in effect_ctns
                                             if c.find(Q("childTnLst")) is None
                                             or c.find("./%s/%s" % (Q("childTnLst"), Q("set"))) is None],
                        presets_without_animEffect=[c.get("id") for c in effect_ctns
                                                    if c.find(Q("childTnLst")) is None
                                                    or c.find("./%s/%s" % (Q("childTnLst"), Q("animEffect"))) is None],
                    ))
            else:
                w("  !! slide %d has p:timing but mainSeq path not found" % idx)

        # effect-node structural checks over whole timing
        bad_class = []
        no_set = []
        no_anim = []
        for c in root.iter(Q("cTn")):
            if c.get("presetClass") is None and c.get("presetID") is None:
                continue
            if c.get("presetClass") != "entr":
                bad_class.append((c.get("id"), c.get("presetClass")))
            ch = c.find(Q("childTnLst"))
            if ch is None or ch.find(Q("set")) is None:
                no_set.append(c.get("id"))
            if ch is None or ch.find(Q("animEffect")) is None:
                no_anim.append(c.get("id"))

        all_spids = set()
        for sp in root.iter(Q("spTgt")):
            all_spids.add(sp.get("spid"))
        missing_targets = sorted(s for s in all_spids if s not in idname)

        # group membership of targets
        grouped_targets = []
        for grp in root.iter(Q("grpSp")):
            gid = None
            gcn = grp.find(Q("nvGrpSpPr") + "/" + Q("cNvPr"))
            if gcn is not None:
                gid = gcn.get("id")
            inner = {c.get("id") for c in grp.iter(Q("cNvPr"))}
            hit = inner & all_spids
            if hit:
                grouped_targets.append((gid, idname.get(gid), sorted(hit)))

        rows.append(dict(idx=idx, part=n, clicks=clicks, order=order, trs=len(trs),
                         tr_children=tr_children, tr_attrs=tr_attrs,
                         timing=tim is not None, has_mainseq=tim is not None and 'mainseq' in dir(),
                         dup_ctn=dup_ctn,
                         bad_class=bad_class, no_set=no_set, no_anim=no_anim,
                         missing=missing_targets, grouped=grouped_targets,
                         n_cnvpr=len(idname), raw_len=len(raw)))
        detail[idx] = dict(idname=idname, root=root)

    total = sum(len(r["clicks"]) for r in rows)
    w("  slides=%d  TOTAL top-level click groups=%d" % (len(rows), total))

    # optional per-slide shape geometry from python-pptx for footer identification
    geo = {}
    if prs is not None:
        for i, sl in enumerate(prs.slides, start=1):
            m = {}
            for sh in sl.shapes:
                try:
                    m[str(sh.shape_id)] = (sh.name, sh.top, sh.left, sh.width, sh.height, str(sh.shape_type))
                except Exception:
                    pass
            geo[i] = m

    for r in rows:
        i = r["idx"]
        w("  --- slide %2d | clicks=%d | transition=%d %s%s | timing=%s | shapes(cNvPr)=%d | xml=%dB"
          % (i, len(r["clicks"]), r["trs"], r["tr_children"], r["tr_attrs"],
             r["timing"], r["n_cnvpr"], r["raw_len"]))
        w("        sld child order: %s" % (r["order"],))
        if r["dup_ctn"]:
            w("        !! DUPLICATE p:cTn ids: %s" % (r["dup_ctn"],))
        if r["bad_class"]:
            w("        !! presetClass != entr: %s" % (r["bad_class"],))
        if r["no_set"]:
            w("        !! effect cTn without p:set: %s" % (r["no_set"],))
        if r["no_anim"]:
            w("        !! effect cTn without p:animEffect: %s" % (r["no_anim"],))
        if r["missing"]:
            w("        !! spTgt spid NOT in slide cNvPr: %s" % (r["missing"],))
        if r["grouped"]:
            w("        grouped-shape targets: %s" % (r["grouped"],))
        for gi, c in enumerate(r["clicks"], start=1):
            w("        click#%d: cond=%s nodeTypes=%s effects=%d distinctTargets=%d presetClass=%s presetID=%s"
              % (gi, c["cond"], c["node_types"], c["n_effects"], len(c["distinct"]),
                 c["preset_classes"], c["preset_ids"]))
            if c["presets_without_set"]:
                w("            !! missing p:set in effect cTn ids %s" % (c["presets_without_set"],))
            if c["presets_without_animEffect"]:
                w("            !! missing p:animEffect in effect cTn ids %s" % (c["presets_without_animEffect"],))
            nm = c["names"]
            w("            targets(%d): %s" % (len(nm), ", ".join(nm[:14]) + (" ..." if len(nm) > 14 else "")))
            if i in geo:
                foot = []
                for s in c["distinct"]:
                    g = geo[i].get(str(s))
                    if g:
                        foot.append("%s(top=%.0f,h=%.0f)" % (g[0], (g[1] or 0) / 914400.0, (g[4] or 0) / 914400.0))
                    else:
                        foot.append("%s(<no geo>)" % s)
                w("            geometry: %s" % (", ".join(foot[:16]),))
    return dict(rows=rows, total=total, worst=worst, zf=zf, path=path, detail=detail, geo=geo)


results = {}
for label, path in FILES:
    results[label] = analyze(path, label)

# ---------------------------------------------------------------- comparisons
def part_map(zf):
    m = {}
    for n in zf.namelist():
        m[n] = sha(zf.read(n))
    return m


w("")
w("=" * 100)
w("COMPARISON: stripped-animation slide XML identity")
for other in ("preanim", "v1perline"):
    a = results["deliverable"]
    b = results[other]
    w("")
    w("  deliverable  vs  %s" % other)
    ma, mb = part_map(a["zf"]), part_map(b["zf"])
    only_a = sorted(set(ma) - set(mb))
    only_b = sorted(set(mb) - set(ma))
    diff = sorted(n for n in set(ma) & set(mb) if ma[n] != mb[n])
    w("    zip entries: deliverable=%d %s=%d ; only_in_deliverable=%s ; only_in_%s=%s"
      % (len(ma), other, len(mb), only_a, other, only_b))
    w("    differing entries (%d): %s" % (len(diff), diff if len(diff) < 40 else diff[:40]))
    # raw byte compare after stripping
    a_named = {n.split("/")[-1]: n for n in slide_parts(a["zf"])}
    b_named = {n.split("/")[-1]: n for n in slide_parts(b["zf"])}
    ident = 0
    nonident = []
    for k in sorted(a_named, key=lambda x: int(re.search(r"(\d+)", x).group(1))):
        if k not in b_named:
            nonident.append((k, "missing in %s" % other))
            continue
        ra, ta, ma_ = strip_anim_bytes(a["zf"].read(a_named[k]))
        rb, tb, mb_ = strip_anim_bytes(b["zf"].read(b_named[k]))
        if ra == rb:
            ident += 1
        else:
            nonident.append((k, "stripped bytes differ (lenA=%d lenB=%d)" % (len(ra), len(rb))))
    w("    slides byte-identical after stripping p:transition+p:timing: %d/%d"
      % (ident, len(a_named)))
    for k, why in nonident:
        w("      !! %s : %s" % (k, why))
    # show a diff snippet for the first non-identical slide
    if nonident and nonident[0][1].startswith("stripped"):
        k = nonident[0][0]
        ra, _, _ = strip_anim_bytes(a["zf"].read(a_named[k]))
        rb, _, _ = strip_anim_bytes(b["zf"].read(b_named[k]))
        sa, sb = ra.decode(), rb.decode()
        j = 0
        while j < min(len(sa), len(sb)) and sa[j] == sb[j]:
            j += 1
        w("      first diff at char %d:" % j)
        w("        A: ...%s..." % sa[max(0, j - 120):j + 200].replace("\n", " "))
        w("        B: ...%s..." % sb[max(0, j - 120):j + 200].replace("\n", " "))
    # notes
    na = {n.split("/")[-1]: sha(a["zf"].read(n)) for n in notes_parts(a["zf"])}
    nb = {n.split("/")[-1]: sha(b["zf"].read(n)) for n in notes_parts(b["zf"])}
    same = [k for k in na if nb.get(k) == na[k]]
    w("    notesSlides: deliverable=%d %s=%d identical=%d  differ=%s"
      % (len(na), other, len(nb), len(same), sorted(set(na) ^ set(nb)) or
         sorted(k for k in na if nb.get(k) != na[k])))

w("")
w("=" * 100)
w("SUMMARY TABLE (clicks per slide)")
for label in ("deliverable", "preanim", "v1perline"):
    r = results[label]
    w("  %-12s total=%d  per-slide=%s" % (label, r["total"],
                                          [len(x["clicks"]) for x in r["rows"]]))

dg = results["deliverable"]
over3 = [x["idx"] for x in dg["rows"] if len(x["clicks"]) > 3]
multi = []
for x in dg["rows"]:
    for gi, c in enumerate(x["clicks"], 1):
        n = c["node_types"].get("clickEffect", 0)
        if n > 1:
            multi.append((x["idx"], gi, n))
w("  deliverable: slides with >3 clicks = %s" % (over3 or "NONE",))
w("  deliverable: click groups containing >1 nodeType=clickEffect = %s" % (multi or "NONE",))
w("  deliverable: max XML depth overall = %s" % (dg["worst"][0],))
w("  deliverable: any part depth>64 = %s" % ([t for t in dg["worst"] if t[0] > 64] or "NONE",))
w("  deliverable: slide2 has timing=%s transition=%s" %
  (dg["rows"][1]["timing"], dg["rows"][1]["tr_children"]))

# ------------------------------------------------- timing-chain / raw structure
def raw_slide(label, n):
    return results[label]["zf"].read("ppt/slides/slide%d.xml" % n)


w("")
w("=" * 100)
w("RAW TIMING STRUCTURE (trimmed, for eyeballing trigger semantics)")
for label, n, groups in (("deliverable", 1, 2), ("deliverable", 3, 1)):
    root = etree.fromstring(raw_slide(label, n))
    tim = root.find(Q("timing"))
    ms = tim.find("./%s/%s/%s/%s/%s/%s/%s" % (Q("tnLst"), Q("par"), Q("cTn"),
                                              Q("childTnLst"), Q("seq"), Q("cTn"),
                                              Q("childTnLst")))
    w("  [%s slide %d] p:seq attrs=%s" % (label, n, {k.split('}')[-1]: v for k, v in ms.getparent().attrib.items()}))
    for gi, gp in enumerate(ms.findall(Q("par"))[:groups], 1):
        s = etree.tostring(gp, pretty_print=True).decode()
        lines = s.splitlines()
        w("  --- group %d of slide %d : %d xml lines ; first 60 lines + effect-node digest:" % (gi, n, len(lines)))
        for ln in lines[:60]:
            w("      " + ln)
        if len(lines) > 60:
            w("      ... (%d more lines)" % (len(lines) - 60))
        seq = []
        for c in gp.iter(Q("cTn")):
            if c.get("presetClass") or c.get("presetID") or c.get("nodeType"):
                st = c.find("./%s/%s" % (Q("stCondLst"), Q("cond")))
                seq.append((c.get("id"), c.get("nodeType"), c.get("presetClass"), c.get("presetID"),
                            c.get("dur"), None if st is None else (st.get("delay"), st.get("evt"))))
        w("      cTn digest (id, nodeType, presetClass, presetID, dur, firstCond) count=%d:" % len(seq))
        for t in seq[:8]:
            w("        %s" % (t,))
        if len(seq) > 8:
            w("        ... and %d more" % (len(seq) - 8))

w("")
w("=" * 100)
w("CLICK-GROUP CHAIN TIMING (does 'one click = everything at once' hold in time?)")
for r in dg["rows"]:
    i = r["idx"]
    if not r["clicks"]:
        continue
    root = results["deliverable"]["detail"][i]["root"]
    tim = root.find(Q("timing"))
    ms = tim.find("./%s/%s/%s/%s/%s/%s/%s" % (Q("tnLst"), Q("par"), Q("cTn"),
                                              Q("childTnLst"), Q("seq"), Q("cTn"), Q("childTnLst")))
    parts = []
    for gi, gp in enumerate(ms.findall(Q("par")), 1):
        durs = []
        for c in gp.iter(Q("cTn")):
            if c.get("presetClass") or c.get("presetID"):
                try:
                    durs.append(int(c.get("dur") or 0))
                except ValueError:
                    durs.append(-1)
        nt = {}
        for c in gp.iter(Q("cTn")):
            if c.get("nodeType"):
                nt[c.get("nodeType")] = nt.get(c.get("nodeType"), 0) + 1
        if durs:
            parts.append("click#%d effects=%d sum(dur)=%dms max(dur)=%dms min=%dms nodeTypes=%s"
                         % (gi, len(durs), sum(durs), max(durs), min(durs), nt))
    w("  slide %2d: %s" % (i, " | ".join(parts)))

# ------------------------------------------------- zip entry diff matrix
w("")
w("=" * 100)
w("ZIP-ENTRY DIFF MATRIX (sha256 per entry) -- which files are actually distinct?")
EXTRA = [
    ("rebuild_a", os.path.join(BASE, "_work", "rebuild_a.pptx")),
    ("rebuild_b", os.path.join(BASE, "_work", "rebuild_b.pptx")),
    ("anim_selftest", os.path.join(BASE, "_work", "anim_selftest.pptx")),
]
sets = {}
for label, path in FILES + [(l, p) for l, p in EXTRA]:
    if not os.path.exists(path):
        continue
    z = zipfile.ZipFile(path)
    sets[label] = (part_map(z), z)
labels = [l for l, _ in FILES + EXTRA if l in sets]
w("  full-file sha256:")
for l in labels:
    p = dict(FILES + EXTRA)[l]
    w("    %-14s %s  %d bytes" % (l, sha(open(p, "rb").read()), os.path.getsize(p)))
w("")
w("  pair-wise count of differing zip entries:")
for x in range(len(labels)):
    for y in range(x + 1, len(labels)):
        a, b = labels[x], labels[y]
        ma, mb = sets[a][0], sets[b][0]
        diff = sorted(k for k in set(ma) & set(mb) if ma[k] != mb[k])
        w("    %-14s vs %-14s : common=%d differing=%d %s" % (a, b, len(set(ma) & set(mb)), len(diff), diff[:8]))
w("")
w("  per-entry sha of the interesting parts:")
interesting = ["ppt/slides/slide1.xml", "ppt/slides/slide3.xml", "ppt/slides/slide7.xml",
               "ppt/slides/slide15.xml", "ppt/presentation.xml", "docProps/app.xml",
               "docProps/core.xml", "ppt/notesSlides/notesSlide3.xml"]
w("    %-34s %s" % ("entry", "  ".join("%-12s" % l for l in labels)))
for e in interesting:
    row = []
    for l in labels:
        h = sets[l][0].get(e)
        row.append("%-12s" % (h[:10] if h else "ABSENT"))
    w("    %-34s %s" % (e, "  ".join(row)))

with open(os.path.join(BASE, "_work", "_audit_anim_verify_output.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(out))
