# -*- coding: utf-8 -*-
"""diag_order.py —— 打印最终 PPTX 每页的标题/首行文字，核对页序。"""
import pathlib
import sys

from pptx import Presentation

p = pathlib.Path(sys.argv[1] if len(sys.argv) > 1
                 else "桥体卫士答辩PPT_v2/桥体卫士答辩PPT_12分钟版.pptx")
prs = Presentation(str(p))
print(f"{p.name}: {len(prs.slides)} 页\n")
for i, s in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text.strip().replace("\n", " / ")
             for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    pics = sum(1 for sh in s.shapes if "PICTURE" in str(sh.shape_type))
    kinds = {}
    for sh in s.shapes:
        k = str(sh.shape_type).split(".")[0].split(" ")[0]
        kinds[k] = kinds.get(k, 0) + 1
    head = texts[0][:38] if texts else "（无文字）"
    print(f"  {i:02d}  text={len(texts):2d} pic={pics}  {head}")
    if i <= 3 or texts[:2]:
        for t in texts[:3]:
            print(f"        · {t[:60]}")
